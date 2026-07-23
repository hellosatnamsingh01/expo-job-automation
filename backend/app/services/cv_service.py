import os
import tempfile
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch


def extract_cv_text(file_path: str) -> str:
    if file_path.lower().endswith(".pdf"):
        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            return ""
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def update_cv_docx(original_path: str, updated_text: str) -> str:
    """Write updated text back into a copy of the Word doc and return the path."""
    doc = Document(original_path)

    # Replace paragraph text while preserving structure
    updated_lines = updated_text.split("\n")
    line_idx = 0
    for para in doc.paragraphs:
        if line_idx < len(updated_lines):
            if para.text.strip():
                para.clear()
                para.add_run(updated_lines[line_idx])
                line_idx += 1

    tmp_docx = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    doc.save(tmp_docx.name)
    return tmp_docx.name


def convert_docx_to_pdf(docx_path: str, output_dir: str = None) -> str:
    """Convert Word doc to PDF using reportlab (basic text extraction)."""
    doc = Document(docx_path)
    text_content = "\n".join(p.text for p in doc.paragraphs)

    if not output_dir:
        output_dir = tempfile.gettempdir()

    pdf_path = os.path.join(output_dir, os.path.basename(docx_path).replace(".docx", ".pdf"))

    styles = getSampleStyleSheet()
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                                rightMargin=inch * 0.75, leftMargin=inch * 0.75,
                                topMargin=inch * 0.75, bottomMargin=inch * 0.75)
    story = []
    for line in text_content.split("\n"):
        if line.strip():
            story.append(Paragraph(line, styles["Normal"]))
            story.append(Spacer(1, 4))
        else:
            story.append(Spacer(1, 8))

    pdf_doc.build(story)
    return pdf_path


def build_pdf_from_text(text: str, output_path: str = None) -> str:
    """Build a clean PDF from plain text. Saves to output_path if given, else a temp file."""
    styles = getSampleStyleSheet()
    if output_path:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        pdf_path = output_path
    else:
        tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
        pdf_path = tmp.name
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                                rightMargin=inch * 0.75, leftMargin=inch * 0.75,
                                topMargin=inch * 0.75, bottomMargin=inch * 0.75)
    story = []
    for line in text.split("\n"):
        if line.strip():
            story.append(Paragraph(line, styles["Normal"]))
            story.append(Spacer(1, 4))
        else:
            story.append(Spacer(1, 8))
    pdf_doc.build(story)
    return pdf_path


def prepare_cv_for_job(original_cv_path: str, updated_text: str, output_path: str = None) -> str:
    """Returns path to a tailored PDF ready to be emailed.
    If output_path is given the file is saved there permanently; otherwise a temp file is used."""
    if original_cv_path.lower().endswith(".pdf"):
        return build_pdf_from_text(updated_text, output_path=output_path)
    tmp_docx = update_cv_docx(original_cv_path, updated_text)
    if output_path:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        pdf_path = convert_docx_to_pdf(tmp_docx, output_dir=os.path.dirname(output_path))
        # rename to the exact requested output_path
        if pdf_path != output_path:
            os.rename(pdf_path, output_path)
        pdf_path = output_path
    else:
        pdf_path = convert_docx_to_pdf(tmp_docx)
    os.unlink(tmp_docx)
    return pdf_path
